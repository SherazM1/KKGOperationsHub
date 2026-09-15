"""Regression coverage for the audited two-row operational load sheet."""

from io import BytesIO

from openpyxl import Workbook
import pytest

from app.services.bol_standard_parser import parse_standard_bol_excel
from app.services.bol_standard_mapper import map_standard_rows_to_records


def _split_workbook(*, dc_name="Regional DC 0608", explicit_skids=None):
    headers = ["KK Load", "TRACKERS", "Carrier", "load#", "KK PO#", "BOL #", "ship date", "DC Name", "DC ADDRESS", "DC CITY", "DC", "DC", "RETAILER PO #", "MABD", "ITEM #", "UPC", "Pallet", "WM WEEK", "Units", "DEPT", "Equipment", "PLT Weight", "Weight", "Value", "Load", 0.03, "Rate", "Accessorials", "All-In Rate", "Delivery Appt Date", "Delivery Appt Time", "Delivery Appt #", "NOTES", "LINE", None]
    continuation = [None] * len(headers)
    for index, value in {10: "ST", 11: "ZIP", 16: "Description", 23: "Each", 24: "Value", 25: "Charge Back"}.items():
        continuation[index] = value
    data = [1, "TRACK-1", "Carrier", "PRO-1", "KK-1", "BOL-1", "2026-09-25", dc_name, "123 Street", "City", "TX", "012345678", "PO-1", "2026-09-26", "ITEM-1", "000123456789", "Test kit", "WEEK36", 71, 13, "DRY VAN", 176, 12496, 248.6, 17650.6, 529.518, 544.32, None, 544.32, "2026-09-26", "0700", "APPT-1", "Keep this note", 1, None]
    if explicit_skids is not None:
        headers.append("PLT QTY")
        continuation.append(None)
        data.append(explicit_skids)
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(headers)
    sheet.append(continuation)
    sheet.append(data)
    sheet.append([None] * len(headers))
    footer = [None] * len(headers)
    footer[26] = 544.32
    sheet.append(footer)
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def test_split_headers_keep_all_columns_and_exclude_footer():
    rows = parse_standard_bol_excel(_split_workbook())
    assert len(rows) == 1
    row = rows[0]
    assert row.source_row_number == 3
    assert row.dc_number == "0608"
    assert row.dc_city_state_zip == "City, TX 012345678"
    assert row.item_description == "Test kit"
    assert (row.unit_qty, row.weight_each, row.total_weight) == ("71", "176", "12496")
    assert row.plt_qty == ""
    assert row.carrier_pro_number == "PRO-1"
    assert row.wm_po == "PO-1"
    assert row.pickup_number == "APPT-1"
    assert len(row.source_values) == 34
    assert row.source_values["Load Value"] == "17650.6"
    assert row.source_values["0.03 Charge Back"] == "529.518"
    assert row.source_values["NOTES"] == "Keep this note"
    record = map_standard_rows_to_records(rows)[0]
    assert record.status == "Ready (pallet count needs review)"
    assert record.is_ready
    assert any("No separate pallet/skid count supplied" in warning for warning in record.warnings)


def test_explicit_skid_count_is_used_and_record_is_ready():
    rows = parse_standard_bol_excel(_split_workbook(explicit_skids=5))
    record = map_standard_rows_to_records(rows)[0]
    assert record.is_ready
    assert record.item_lines[0].skids == "5"
    assert record.item_lines[0].pallet_qty == "71"


@pytest.mark.parametrize("name", ["Regional warehouse 0608", "DC 0608 / DC 0609"])
def test_dc_number_is_not_guessed(name):
    rows = parse_standard_bol_excel(_split_workbook(dc_name=name, explicit_skids=5))
    assert rows[0].dc_number == ""
    assert "DC #" in map_standard_rows_to_records(rows)[0].missing_required_fields


def test_split_csv_uses_same_headers_and_source_row_numbers():
    from openpyxl import load_workbook
    import csv
    from io import StringIO

    sheet = load_workbook(_split_workbook(), data_only=True).active
    text = StringIO()
    csv.writer(text).writerows(sheet.values)
    source = BytesIO(text.getvalue().encode())
    source.name = "loads.csv"
    rows = parse_standard_bol_excel(source)
    assert len(rows) == 1
    assert rows[0].source_row_number == 3
    assert rows[0].dc_city_state_zip == "City, TX 012345678"
    assert rows[0].source_values["Load Value"] == "17650.6"
